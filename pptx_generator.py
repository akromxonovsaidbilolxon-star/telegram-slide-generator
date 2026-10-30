import os
import io
import urllib.parse
import requests
from PIL import Image, ImageDraw, ImageFont
from google import genai
from google.genai import types
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pydantic import BaseModel, Field

# ----------------- THEMES & PALETTES -----------------
THEMES = {
    "obsidian": {
        "name": "Obsidian Cyan",
        "dark_bg": RGBColor(10, 15, 29),
        "card_bg": RGBColor(19, 29, 53),
        "card_border": RGBColor(45, 65, 110),
        "text_main": RGBColor(255, 255, 255),
        "text_muted": RGBColor(180, 195, 220),
        "accent": RGBColor(56, 189, 248),
        "accent_badge": RGBColor(14, 116, 144)
    },
    "emerald": {
        "name": "Nordic Emerald",
        "dark_bg": RGBColor(6, 20, 24),
        "card_bg": RGBColor(13, 37, 44),
        "card_border": RGBColor(28, 75, 88),
        "text_main": RGBColor(255, 255, 255),
        "text_muted": RGBColor(165, 190, 200),
        "accent": RGBColor(45, 212, 191),
        "accent_badge": RGBColor(15, 118, 110)
    },
    "purple": {
        "name": "Deep Royal Purple",
        "dark_bg": RGBColor(15, 12, 34),
        "card_bg": RGBColor(30, 24, 66),
        "card_border": RGBColor(75, 55, 145),
        "text_main": RGBColor(255, 255, 255),
        "text_muted": RGBColor(195, 185, 235),
        "accent": RGBColor(192, 132, 252),
        "accent_badge": RGBColor(107, 33, 168)
    }
}

# ----------------- SCHEMAS -----------------
class MetricBox(BaseModel):
    value: str = Field(description="Display stat, e.g. '85%', '10x', '99.4%'")
    title: str = Field(description="Short label, max 3 words")
    desc: str = Field(description="Brief 1-sentence explanation")

class PillarCard(BaseModel):
    number: str = Field(description="'01', '02', '03'")
    title: str = Field(description="Headline of pillar")
    desc: str = Field(description="Two clear sentences")

class SlideContent(BaseModel):
    slide_number: int
    layout_type: str = Field(description="'title', 'cards_grid', 'stats_row', 'split_image', 'closing'")
    eyebrow: str
    title: str
    subtitle: str | None = None
    source_citation: str | None = None
    bullet_points: list[str] = Field(default=[])
    cards: list[PillarCard] = Field(default=[])
    metrics: list[MetricBox] = Field(default=[])
    takeaway_callout: str | None = None
    ai_image_prompt: str | None = Field(default=None, description="2-3 specific English keywords, e.g. 'economics stock market', 'factory production line'")

class MasterpiecePresentation(BaseModel):
    topic: str
    language: str
    theme_key: str = "obsidian"
    total_slides: int
    slides: list[SlideContent]

class OutlineProposal(BaseModel):
    variant_a_title: str
    variant_a_points: list[str]
    variant_b_title: str
    variant_b_points: list[str]

# ----------------- AI BRAINSTORMING HELPER -----------------
def generate_outline_options(topic: str, slide_count: int, audience_type: str, language: str = "uz") -> OutlineProposal:
    client = genai.Client(api_key=os.environ.get("GEMINI_API_KEY"))

    prompt = (
        f"Mavzu: '{topic}'. Slaydlar soni: {slide_count}. Yo'nalish: {audience_type}.\n"
        f"Foydalanuvchiga aynan {slide_count} ta slayddan iborat 2 xil professional reja (Variant A va Variant B) taklif qil.\n"
        f"Har ikkala variantda aniq {slide_count} ta band bo'lsin. Til: o'zbek tili."
    )

    res = client.models.generate_content(
        model="gemini-3.8-flash",
        contents=prompt,
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=OutlineProposal,
            temperature=0.3
        )
    )
    return res.parsed

# ----------------- SMART IMAGE FETCHER (TOPIC-ACCURATE) -----------------
def fetch_image_secure(prompt: str | None, custom_bytes: bytes | None = None) -> io.BytesIO | None:
    # 1. Prioritize user's own uploaded picture if available
    if custom_bytes:
        try:
            return io.BytesIO(custom_bytes)
        except Exception:
            pass

    if not prompt:
        prompt = "business finance technology"

    clean_q = " ".join(prompt.strip().replace("-", " ").split()[:3])

    # 2. Try Google Custom Search (Image)
    api_key = os.environ.get("GOOGLE_SEARCH_API_KEY")
    cx = os.environ.get("GOOGLE_CSE_ID")
    if api_key and cx:
        try:
            url = "https://www.googleapis.com/customsearch/v1"
            params = {
                "key": api_key,
                "cx": cx,
                "q": clean_q,
                "searchType": "image",
                "num": 2,
                "imgSize": "large",
                "safe": "active"
            }
            res = requests.get(url, params=params, timeout=6)
            if res.status_code == 200:
                items = res.json().get("items", [])
                for item in items:
                    l = item.get("link")
                    if l:
                        ir = requests.get(l, headers={"User-Agent": "Mozilla/5.0"}, timeout=5)
                        if ir.status_code == 200 and len(ir.content) > 4000:
                            bio = io.BytesIO(ir.content)
                            img = Image.open(bio)
                            img.verify()
                            bio.seek(0)
                            return bio
        except Exception:
            pass

    # 3. Try Pixabay
    pix_key = os.environ.get("PIXABAY_API_KEY")
    if pix_key:
        try:
            p_res = requests.get("https://pixabay.com/api/", params={
                "key": pix_key,
                "q": clean_q,
                "image_type": "photo",
                "orientation": "horizontal",
                "per_page": 3
            }, timeout=5)
            if p_res.status_code == 200:
                hits = p_res.json().get("hits", [])
                if hits:
                    u = hits[0].get("webformatURL")
                    ir = requests.get(u, timeout=5)
                    if ir.status_code == 200:
                        return io.BytesIO(ir.content)
        except Exception:
            pass

    # 4. Fallback: Fast AI Generation with Pollinations Flux
    try:
        encoded = urllib.parse.quote_plus(f"{clean_q} cinematic professional 8k")
        url = f"https://image.pollinations.ai/prompt/{encoded}?width=1024&height=640&model=flux&nologo=true"
        resp = requests.get(url, timeout=7)
        if resp.status_code == 200 and len(resp.content) > 4000:
            bio = io.BytesIO(resp.content)
            img = Image.open(bio)
            img.verify()
            bio.seek(0)
            return bio
    except Exception:
        pass

    return None

# ----------------- CONTENT GENERATION VIA GEMINI -----------------
def generate_presentation_json(
    topic: str,
    slide_count: int = 10,
    language: str = "uz",
    user_notes: str | None = None,
    theme_key: str = "obsidian"
) -> MasterpiecePresentation:
    client = genai.Client(api_key=os.environ.get("GEMINI_API_KEY"))

    lang_map = {
        "uz": "O'zbek tili (barcha matnlar 100% ravon o'zbek tilida)",
        "ru": "Русский язык (весь контент строго на русском)",
        "en": "English language (executive presentation grade)"
    }
    target_language = lang_map.get(language, "O'zbek tili")

    research_context = ""
    try:
        research_prompt = (
            f"Search Google for real verified facts, statistics, definitions, and core insights on: '{topic}'. "
            f"User notes or plan: {user_notes or 'None'}"
        )
        search_res = client.models.generate_content(
            model="gemini-3.8-flash",
            contents=research_prompt,
            config=types.GenerateContentConfig(
                tools=[types.Tool(google_search=types.GoogleSearch())],
                temperature=0.2
            )
        )
        if search_res and search_res.text:
            research_context = search_res.text
    except Exception as e:
        print(f"[Search Warning]: {e}")

    # Check if user passed a numbered list (Reja)
    outline_lines = [l for l in (user_notes or "").split("\n") if any(l.strip().startswith(f"{i}.") for i in range(1, 30))]
    target_count = len(outline_lines) if len(outline_lines) >= 4 else slide_count

    system_instruction = (
        f"You are an elite slide designer. Generate EXACTLY {target_count} slides in the 'slides' array. "
        "CRITICAL RULES TO AVOID EMPTY CARDS:\n"
        "1. For 'cards_grid' layout, you MUST ALWAYS provide EXACTLY 3 cards inside the 'cards' list. Never leave it empty or with only 1 card.\n"
        "2. For 'stats_row' layout, you MUST ALWAYS provide EXACTLY 3 metrics inside the 'metrics' list.\n"
        "3. For 'split_image' layout, provide 3 to 4 detailed bullet points, AND provide an English 'ai_image_prompt' directly relevant to the slide topic (e.g. 'stock market exchange board', 'industrial automation robot', 'central bank building').\n"
        "4. Slide 1 MUST be 'title', and Slide {target_count} MUST be 'closing'.\n"
        "5. Keep headlines short (under 35 chars) so they don't break on mobile screens."
    )

    synthesis_prompt = (
        f"Topic: {topic}\n"
        f"Exact Slides Count: {target_count}\n"
        f"Language: {target_language}\n"
        f"User Provided Plan / Custom Outline:\n{user_notes or 'None'}\n\n"
        f"Research Context:\n{research_context[:3000]}"
    )

    response = client.models.generate_content(
        model="gemini-3.8-flash",
        contents=synthesis_prompt,
        config=types.GenerateContentConfig(
            system_instruction=system_instruction,
            response_mime_type="application/json",
            response_schema=MasterpiecePresentation,
            temperature=0.2
        ),
    )
    result = response.parsed
    result.theme_key = theme_key
    return result

# ----------------- POWERPOINT BUILDER -----------------
def create_pptx_file(
    plan: MasterpiecePresentation,
    output_path: str = "output.pptx",
    custom_img_bytes: bytes | None = None
) -> str:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    theme = THEMES.get(plan.theme_key, THEMES["obsidian"])
    blank_layout = prs.slide_layouts[6]
    custom_image_used = False

    for s in plan.slides:
        slide = prs.slides.add_slide(blank_layout)

        # Full dark background
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = theme["dark_bg"]
        bg.line.fill.background()

        # ---------------- 1. TITLE SLIDE ----------------
        if s.layout_type == "title":
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.2), Inches(0.12), Inches(3.0))
            bar.fill.solid()
            bar.fill.fore_color.rgb = theme["accent"]
            bar.line.fill.background()

            tb = slide.shapes.add_textbox(Inches(1.4), Inches(2.0), Inches(10.5), Inches(3.6))
            tf = tb.text_frame
            tf.word_wrap = True

            pe = tf.paragraphs[0]
            pe.text = s.eyebrow.upper()
            pe.font.size = Pt(14)
            pe.font.bold = True
            pe.font.color.rgb = theme["accent"]
            pe.space_after = Pt(12)

            pt = tf.add_paragraph()
            pt.text = s.title
            pt.font.size = Pt(44)
            pt.font.bold = True
            pt.font.color.rgb = theme["text_main"]
            pt.space_after = Pt(16)

            if s.subtitle:
                ps = tf.add_paragraph()
                ps.text = s.subtitle
                ps.font.size = Pt(20)
                ps.font.color.rgb = theme["text_muted"]
            continue

        # ---------------- HEADER FOR ALL OTHER SLIDES ----------------
        hb = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.0), Inches(1.1))
        hf = hb.text_frame
        hf.word_wrap = True

        he = hf.paragraphs[0]
        he.text = s.eyebrow.upper()
        he.font.size = Pt(12)
        he.font.bold = True
        he.font.color.rgb = theme["accent"]
        he.space_after = Pt(4)

        ht = hf.add_paragraph()
        ht.text = s.title
        ht.font.size = Pt(28)
        ht.font.bold = True
        ht.font.color.rgb = theme["text_main"]

        # Number badge on top right
        nb = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.7), Inches(0.5), Inches(0.8), Inches(0.42))
        nb.fill.solid()
        nb.fill.fore_color.rgb = theme["card_bg"]
        nb.line.color.rgb = theme["accent"]
        np = nb.text_frame.paragraphs[0]
        np.text = f"{s.slide_number:02d}"
        np.font.size = Pt(13)
        np.font.bold = True
        np.font.color.rgb = theme["accent"]
        np.alignment = PP_ALIGN.CENTER

        # ---------------- 2. CARDS GRID (NO EMPTY BOXES) ----------------
        if s.layout_type == "cards_grid":
            items = list(s.cards) if s.cards else []
            # Fill fallback items if Gemini provided fewer than 3
            default_labels = ["Asosiy Tamoyil", "Tahliliy Yondashuv", "Amaliy Natija"]
            while len(items) < 3:
                idx = len(items)
                items.append(PillarCard(
                    number=f"0{idx+1}",
                    title=default_labels[idx],
                    desc=s.bullet_points[idx] if idx < len(s.bullet_points) else "Ushbu yo'nalish bo'yicha tizimli yondashuv va tahliliy tadqiqotlar asosiy mezon hisoblanadi."
                ))

            c_w = Inches(3.64)
            c_h = Inches(4.3)
            gap = Inches(0.4)

            for i, c in enumerate(items[:3]):
                cx = Inches(0.8) + i * (c_w + gap)
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(1.9), c_w, c_h)
                card.fill.solid()
                card.fill.fore_color.rgb = theme["card_bg"]
                card.line.color.rgb = theme["card_border"]

                ctb = slide.shapes.add_textbox(cx + Inches(0.3), Inches(2.1), c_w - Inches(0.6), c_h - Inches(0.4))
                ctf = ctb.text_frame
                ctf.word_wrap = True

                p0 = ctf.paragraphs[0]
                p0.text = c.number
                p0.font.size = Pt(14)
                p0.font.bold = True
                p0.font.color.rgb = theme["accent"]
                p0.space_after = Pt(8)

                p1 = ctf.add_paragraph()
                p1.text = c.title
                p1.font.size = Pt(19)
                p1.font.bold = True
                p1.font.color.rgb = theme["text_main"]
                p1.space_after = Pt(12)

                p2 = ctf.add_paragraph()
                p2.text = c.desc
                p2.font.size = Pt(14)
                p2.font.color.rgb = theme["text_muted"]

        # ---------------- 3. STATS ROW (NO EMPTY BOXES) ----------------
        elif s.layout_type == "stats_row":
            m_items = list(s.metrics) if s.metrics else []
            default_metrics = [
                MetricBox(value="85%", title="SAMARADORLIK", desc="Tizimli optimallashuv darajasi"),
                MetricBox(value="10x", title="TEZLIK", desc="Operatsion boshqaruv jadalligi"),
                MetricBox(value="100%", title="KAFOLAT", desc="Rejalashtirilgan natijaga erishish")
            ]
            while len(m_items) < 3:
                m_items.append(default_metrics[len(m_items)])

            m_w = Inches(3.64)
            m_h = Inches(4.3)
            gap = Inches(0.4)

            for i, m in enumerate(m_items[:3]):
                mx = Inches(0.8) + i * (m_w + gap)
                box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, mx, Inches(1.9), m_w, m_h)
                box.fill.solid()
                box.fill.fore_color.rgb = theme["card_bg"]
                box.line.color.rgb = theme["card_border"]

                mtb = slide.shapes.add_textbox(mx + Inches(0.3), Inches(2.3), m_w - Inches(0.6), m_h - Inches(0.6))
                mtf = mtb.text_frame
                mtf.word_wrap = True

                pv = mtf.paragraphs[0]
                pv.text = m.value
                pv.font.size = Pt(46)
                pv.font.bold = True
                pv.font.color.rgb = theme["accent"] if i == 0 else theme["text_main"]
                pv.space_after = Pt(10)

                ptit = mtf.add_paragraph()
                ptit.text = m.title.upper()
                ptit.font.size = Pt(16)
                ptit.font.bold = True
                ptit.font.color.rgb = theme["text_main"]
                ptit.space_after = Pt(10)

                pd = mtf.add_paragraph()
                pd.text = m.desc
                pd.font.size = Pt(14)
                pd.font.color.rgb = theme["text_muted"]

        # ---------------- 4. SPLIT IMAGE LAYOUT (FIXED: NO OVERLAPPING DARK BOX) ----------------
        elif s.layout_type == "split_image":
            # Use user uploaded photo on the first visual slide, then fallback to API search
            current_bytes = custom_img_bytes if not custom_image_used else None
            img_io = fetch_image_secure(s.ai_image_prompt or s.title, current_bytes)
            if current_bytes and img_io:
                custom_image_used = True

            # If image exists, give text 6 inches; else full width 11.7 inches
            text_width = Inches(5.9) if img_io else Inches(11.7)

            tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.9), text_width, Inches(4.5))
            tf = tb.text_frame
            tf.word_wrap = True

            bullets = s.bullet_points if s.bullet_points else [
                "Strategik tahlil va bozor qonuniyatlariga asoslangan boshqaruv.",
                "Resurslardan samarali foydalanish va rentabellikni oshirish choralari.",
                "Raqobatbardoshlikni oshirishda innovatsion mexanizmlar roli."
            ]

            for idx, pt_text in enumerate(bullets[:4]):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.text = f"•  {pt_text}"
                p.font.size = Pt(16)
                p.font.color.rgb = theme["text_main"]
                p.space_after = Pt(16)

            # Insert Image on the right without ANY overlapping box
            if img_io:
                try:
                    slide.shapes.add_picture(img_io, Inches(7.1), Inches(1.9), width=Inches(5.4), height=Inches(4.3))
                except Exception as e:
                    print(f"Error placing image: {e}")

        # ---------------- 5. CLOSING SLIDE ----------------
        elif s.layout_type == "closing":
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.9), Inches(11.3), Inches(4.3))
            box.fill.solid()
            box.fill.fore_color.rgb = theme["card_bg"]
            box.line.color.rgb = theme["card_border"]

            ctb = slide.shapes.add_textbox(Inches(1.5), Inches(2.4), Inches(10.3), Inches(3.3))
            ctf = ctb.text_frame
            ctf.word_wrap = True

            p1 = ctf.paragraphs[0]
            p1.text = "STRATEGIK XULOSA"
            p1.font.size = Pt(14)
            p1.font.bold = True
            p1.font.color.rgb = theme["accent"]
            p1.space_after = Pt(14)

            p2 = ctf.add_paragraph()
            p2.text = s.title.upper()
            p2.font.size = Pt(32)
            p2.font.bold = True
            p2.font.color.rgb = theme["text_main"]
            p2.space_after = Pt(16)

            if s.takeaway_callout or s.subtitle:
                p3 = ctf.add_paragraph()
                p3.text = s.takeaway_callout or s.subtitle
                p3.font.size = Pt(18)
                p3.font.color.rgb = theme["text_muted"]

        # Footer citation
        if s.source_citation:
            fc = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.4))
            fp = fc.text_frame.paragraphs[0]
            fp.text = s.source_citation
            fp.font.size = Pt(11)
            fp.font.color.rgb = theme["text_muted"]

    prs.save(output_path)
    return output_path

# ----------------- TELEGRAM PREVIEW CARD GENERATOR -----------------
def generate_preview_card(plan: MasterpiecePresentation, output_img_path: str = "preview.png") -> str:
    W, H = 1200, 675
    theme = THEMES.get(plan.theme_key, THEMES["obsidian"])
    
    bg_tuple = (theme["dark_bg"][0], theme["dark_bg"][1], theme["dark_bg"][2])
    card_tuple = (theme["card_bg"][0], theme["card_bg"][1], theme["card_bg"][2])
    border_tuple = (theme["card_border"][0], theme["card_border"][1], theme["card_border"][2])
    accent_tuple = (theme["accent"][0], theme["accent"][1], theme["accent"][2])

    img = Image.new("RGB", (W, H), color=bg_tuple)
    draw = ImageDraw.Draw(img)

    draw.rectangle([(0, 0), (W, 10)], fill=accent_tuple)

    try:
        font_large = ImageFont.truetype("DejaVuSans-Bold.ttf", 40)
        font_sub = ImageFont.truetype("DejaVuSans.ttf", 20)
        font_tag = ImageFont.truetype("DejaVuSans-Bold.ttf", 15)
    except Exception:
        font_large = ImageFont.load_default()
        font_sub = ImageFont.load_default()
        font_tag = ImageFont.load_default()

    draw.text((70, 65), f"AKADEMIX AI  •  {len(plan.slides)} SLIDES READY", fill=accent_tuple, font=font_tag)

    title_text = plan.topic.upper()
    if len(title_text) > 42:
        title_text = title_text[:39] + "..."
    draw.text((70, 105), title_text, fill=(248, 250, 252), font=font_large)

    card_w, card_h = 330, 370
    y_card, gap = 210, 35

    for i, s in enumerate(plan.slides[:3]):
        x_card = 70 + i * (card_w + gap)
        draw.rounded_rectangle([(x_card, y_card), (x_card + card_w, y_card + card_h)], radius=12, fill=card_tuple, outline=border_tuple, width=2)
        draw.rounded_rectangle([(x_card + 20, y_card + 20), (x_card + 75, y_card + 50)], radius=6, fill=accent_tuple)
        draw.text((x_card + 32, y_card + 25), f"{s.slide_number:02d}", fill=(255, 255, 255), font=font_tag)

        s_title = s.title if len(s.title) < 26 else s.title[:23] + "..."
        draw.text((x_card + 20, y_card + 75), s_title, fill=(248, 250, 252), font=font_tag)
        draw.text((x_card + 20, y_card + 115), f"Layout: {s.layout_type}", fill=(165, 180, 205), font=font_sub)

    img.save(output_img_path)
    return output_img_path
